"""Word 文档工具（.docx 走 python-docx；.doc/.wps 旧格式走 WPS COM）。

用法（Windows 一律用 py -3，`python` 可能是 Microsoft Store 别名 stub）：
  py -3 word_tool.py read <file>                      # 提取全文（含表格）
  py -3 word_tool.py info <file>                      # 统计信息
  py -3 word_tool.py create <out.docx> --title X --from-md a.md
  py -3 word_tool.py edit <file> --replace "旧=新" [--replace "a=b"] [--out out.docx]
  py -3 word_tool.py convert <src> <dst>              # 导出 PDF 等（WPS COM）
  py -3 word_tool.py compare A.docx B.docx --out-dir out [--author 名]
                                                      # 比对：diff.txt + diff.html + tracked.docx
  py -3 word_tool.py apply-style <file> [--spec standard] [--out out.docx] [--dry-run]
                                                      # 对已有 docx 套版式（页面/命名样式/字体四属性；
                                                      #  内容零改动断言不通过则拒绝产出、原文件不动）
  py -3 word_tool.py table-style <file> [--spec standard] [--out out.docx] [--dry-run]
                                                      # 表格样式：表头底纹/边框/跨页重复表头/表内字号

参数形态（实测易踩）：read/info/edit 的文件、convert 的 <src> <dst>、compare 的 A/B 全是
**位置参数**；compare 的 --out-dir 是**必填**。

Markdown 支持：标题(#/##/###)、无序列表(-)、有序列表(1.)、表格(| a | b |)。

compare 说明（2026-09-12 新增，据实测）：
  - 红线修订走 WPS COM app.CompareDocuments(Document, Document)——必须传 Document 对象，传路径会
    TypeError；每次比较使用独立 KWPS 实例（同实例第二次比较实测返回 0 条修订），保存前校验
    Revisions.Count > 0，为 0 直接报错而不是产出空文件。
  - 作者名默认取系统用户名，可用 --author 覆盖。实现方式：WPS 的 CompareDocuments 作者名不受
    app.UserName 控制（实测赋值无效，且会持久化污染 WPS 全局配置），Revisions.Author 又是只读；
    故在产物层面把 w:ins/w:del 的 w:author 属性统一改写为指定作者（不动批注作者）。
  - 输入文件若本身带修订（w:ins/w:del），python-docx 读出的正文不可信，默认拒绝比对；
    确认可接受时加 --allow-tracked 继续。
  - 实测局限：红线修订只覆盖正文与表格。页眉/页脚差异不会被 WPS 记入 Revisions，
    "拒绝全部修订"也回不到 A 的页眉（结果里保留 B 的页眉），工具会就此打出告警。
"""
import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import cli_guard  # noqa: E402
import wps_com  # noqa: E402

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

EAST_ASIA_FONT = "微软雅黑"


def _set_east_asia(rpr, font=EAST_ASIA_FONT):
    from docx.oxml.ns import qn

    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def _runs_md(para) -> str:
    """把段落 runs 拼成行内 Markdown（加粗 run 包 `**`）。

    2026-09-17 修：导出侧此前用 `para.text`（纯文本），与生成侧的 `**加粗**` 解析不对称，
    导致「导出 md → 再生成」往返丢失加粗。加粗标记不紧贴空白（与 _MD_BOLD_RE 的 `**非空**` 对齐）；
    表格单元格不加粗（单元格按 ` | ` 切分且生成侧对单元格不做行内解析，写 `**` 会被当字面量）。
    """
    parts = []
    for r in para.runs:
        t = r.text or ""
        if not t:
            continue
        core = t.strip()
        if r.bold and core:
            lead = t[: len(t) - len(t.lstrip())]
            trail = t[len(t.rstrip()) :]
            parts.append(lead + "**" + core + "**" + trail)
        else:
            parts.append(t)
    return "".join(parts)


def read_docx(path):
    from docx import Document

    doc = Document(path)
    out = []
    for para in doc.paragraphs:
        text = _runs_md(para).strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        prefix = {
            "title": "# ",
            "heading 1": "# ",
            "heading 2": "## ",
            "heading 3": "### ",
            "heading 4": "#### ",
            "heading 5": "##### ",
        }.get(style, "")
        out.append(prefix + text)
    for ti, table in enumerate(doc.tables, 1):
        out.append(f"[表格 {ti}]")
        for row in table.rows:
            out.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
    return "\n".join(out)


def cmd_read(args):
    path = args.file
    if path.lower().endswith(".docx"):
        print(read_docx(path))
        return 0
    print(wps_com.read_text(path, "word"))
    return 0


def cmd_info(args):
    from docx import Document

    doc = Document(args.file)
    paras = [p for p in doc.paragraphs if p.text.strip()]
    chars = sum(len(p.text) for p in doc.paragraphs)
    tables = len(doc.tables)
    sections = len(doc.sections)
    print(f"文件: {args.file}")
    print(f"段落数: {len(paras)}  总字符: {chars}  表格数: {tables}  节数: {sections}")
    return 0


def _flush_table(doc, buf):
    if not buf:
        return
    rows = []
    for line in buf:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(width):
            table.cell(ri, ci).text = row[ci] if ci < len(row) else ""
    buf.clear()


# Markdown 行内标记（2026-09-15 修：此前只认到 3 级标题、且 `**加粗**` 会把星号写进文档）
_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _strip_md(text: str) -> str:
    """去掉行内标记（标题里不落 `**`）。"""
    return _MD_BOLD_RE.sub(r"\1", text or "").strip()


def _add_text_runs(para, text: str) -> None:
    """把含 `**加粗**` 的文本写入段落：按标记切 run，星号不落盘，并逐个设 eastAsia。"""
    text = text or ""
    pos = 0
    for m in _MD_BOLD_RE.finditer(text):
        if m.start() > pos:
            r = para.add_run(text[pos:m.start()])
            _set_east_asia(r._element.get_or_add_rPr())
        r = para.add_run(m.group(1))
        r.bold = True
        _set_east_asia(r._element.get_or_add_rPr())
        pos = m.end()
    if pos < len(text):
        r = para.add_run(text[pos:])
        _set_east_asia(r._element.get_or_add_rPr())
    if not text:
        para.add_run("")


def md_to_docx(doc, text):
    from docx.shared import Pt

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    _set_east_asia(normal.element.get_or_add_rPr())

    buf = []
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if re.match(r"^\s*\|", line):
            buf.append(line)
            continue
        _flush_table(doc, buf)
        s = line.strip()
        m = _MD_HEADING_RE.match(s)
        if m:
            # 支持 1–6 级标题（主人 2026-09-15：最多用到 5 级；6 级及以上按模板可容）
            doc.add_heading(_strip_md(m.group(2)), level=min(len(m.group(1)), 6))
        elif re.match(r"^[-*]\s+", s):
            para = doc.add_paragraph(style="List Bullet")
            _add_text_runs(para, re.sub(r"^[-*]\s+", "", s))
        elif re.match(r"^\d+[.、]\s+", s):
            para = doc.add_paragraph(style="List Number")
            _add_text_runs(para, re.sub(r"^\d+[.、]\s+", "", s))
        else:
            para = doc.add_paragraph()
            _add_text_runs(para, s)
    _flush_table(doc, buf)


def cmd_create(args):
    from docx import Document

    doc = Document()
    if args.title:
        doc.add_heading(args.title, level=0)
    body = args.body or ""
    if args.from_md:
        body += "\n" + Path(args.from_md).read_text(encoding="utf-8-sig")
    if body.strip():
        md_to_docx(doc, body)
    doc.save(args.out)
    print(f"OK: 已生成 {args.out}")
    return 0


def _replace_in_para(para, old, new):
    if old not in para.text:
        return False
    new_text = para.text.replace(old, new)
    if para.runs:
        para.runs[0].text = new_text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(new_text)
    return True


def cmd_edit(args):
    from docx import Document

    pairs = [tuple(p.split("=", 1)) for p in args.replace]
    doc = Document(args.file)
    touched = 0
    for para in doc.paragraphs:
        for old, new in pairs:
            if _replace_in_para(para, old, new):
                touched += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for old, new in pairs:
                        if _replace_in_para(para, old, new):
                            touched += 1
    out = args.out or args.file
    doc.save(out)
    print(f"OK: 替换涉及 {touched} 处段落 -> {out}")
    return 0


def cmd_convert(args):
    return wps_com.convert(args.src, args.dst, "word")


# ============================ compare（2026-09-12 新增） ============================

def _system_user():
    import getpass
    import os

    try:
        return getpass.getuser()
    except Exception:
        return os.environ.get("USERNAME") or os.environ.get("USER") or "unknown"


def _tracked_counts(path):
    """返回 (w:ins 数, w:del 数)；非 docx 或读取失败返回 (0, 0)。"""
    import zipfile

    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
    except Exception:
        return (0, 0)
    ins = len(re.findall(r"<w:ins[ >]", xml))
    dele = len(re.findall(r"<w:del[ >]", xml))
    return (ins, dele)


def _docx_units(path):
    """把 docx 摊平成"可比对单元"行：正文段落 + 表格行 + 页眉/页脚文本。"""
    from docx import Document

    doc = Document(path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(f"[正文] {text}")
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append(f"[表格{ti} 第{ri}行] " + " | ".join(cells))
    for si, section in enumerate(doc.sections, 1):
        for kind, part in (("页眉", section.header), ("页脚", section.footer)):
            for para in part.paragraphs:
                text = para.text.strip()
                if text:
                    lines.append(f"[{kind}{si}] {text}")
    return lines


def _write_text_diff(a, b, out_dir, a_lines, b_lines):
    import difflib

    diff = list(
        difflib.unified_diff(
            a_lines, b_lines, fromfile=Path(a).name, tofile=Path(b).name, lineterm="", n=2
        )
    )
    header = [
        f"# 文本差异摘要  A={a}  B={b}",
        f"# A 单元数={len(a_lines)}  B 单元数={len(b_lines)}  差异行数={len(diff)}",
        "# 粒度：段落 / 表格行 / 页眉页脚（- 为 A，+ 为 B）",
        "",
    ]
    (out_dir / "diff.txt").write_text("\n".join(header + diff) + "\n", encoding="utf-8")
    html = difflib.HtmlDiff().make_file(
        a_lines, b_lines, fromdesc=Path(a).name, todesc=Path(b).name, context=True, numlines=3
    )
    (out_dir / "diff.html").write_text(html, encoding="utf-8")
    return diff


def _rewrite_revision_author(path, author):
    """把 tracked.docx 里 w:ins/w:del 上的 w:author 统一改写为 author，返回改写处数。

    实测（2026-09-12）：WPS 12.0 的 CompareDocuments 作者名**不随 app.UserName 变化**
    （UserName 已是新值仍写旧显示名），且 Revisions.Author 为只读（赋值报 AttributeError），
    改 app.UserName 还会把新值持久化进 WPS 全局配置（副作用）。因此在 OOXML 层改写属性，
    既确定生效，也不动用户配置。只改 w:ins/w:del 元素，不碰批注等其他作者信息。
    """
    import os
    import zipfile
    from xml.sax.saxutils import escape

    escaped = escape(author, {'"': "&quot;"})
    tmp = str(path) + ".tmp"
    changed = [0]

    def fix_tag(match):
        new_tag, n = re.subn(r'w:author="[^"]*"', f'w:author="{escaped}"', match.group(0))
        changed[0] += n
        return new_tag

    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                try:
                    text = data.decode("utf-8")
                except UnicodeDecodeError:
                    zout.writestr(item, data)
                    continue
                new_text = re.sub(r"<w:(?:ins|del)\b[^>]*>", fix_tag, text)
                if new_text != text:
                    data = new_text.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)
    return changed[0]


def _wps_compare(a, b, out_path):
    """用独立 KWPS 实例做原生比较，返回 (修订条数, 修订摘要列表)。失败抛异常。"""
    import os

    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = None
    doc_a = doc_b = act = None
    try:
        # 每次比较都用新的独立实例：同实例内连续第二次比较实测返回 0 条修订
        app = win32com.client.DispatchEx("KWPS.Application")
        for attr, value in (("Visible", False), ("DisplayAlerts", 0)):
            try:
                setattr(app, attr, value)
            except Exception:
                pass
        # 注意：这里刻意不设置 app.UserName——实测它既改不动修订作者，又会持久化污染 WPS 全局配置

        doc_a = app.Documents.Open(a, ReadOnly=False)
        doc_b = app.Documents.Open(b, ReadOnly=True)
        # 关键：必须传 Document 对象；传路径会 TypeError
        app.CompareDocuments(doc_a, doc_b)
        act = app.ActiveDocument
        count = int(act.Revisions.Count)
        summary = []
        for i in range(1, min(count, 20) + 1):
            try:
                rev = act.Revisions(i)
                summary.append((i, rev.Type, rev.Author, (rev.Range.Text or "")[:60]))
            except Exception:
                break
        if count <= 0:
            raise RuntimeError(
                "WPS 比较返回 0 条修订：两份文档可能无差异，或 WPS 实例复用导致"
                "（本工具每次已用独立实例，可重跑一次确认）"
            )
        if os.path.exists(out_path):
            os.remove(out_path)
        act.SaveAs(out_path)
        return count, summary
    finally:
        for doc in (act, doc_a, doc_b):
            try:
                if doc is not None:
                    doc.Close(False)
            except Exception:
                pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def cmd_compare(args):
    import os

    a, b = os.path.abspath(args.doc_a), os.path.abspath(args.doc_b)
    for path in (a, b):
        if not os.path.isfile(path):
            print(f"错误: 文件不存在 {path}", file=sys.stderr)
            return 2
    if not (a.lower().endswith(".docx") and b.lower().endswith(".docx")):
        print("错误: compare 仅支持 .docx（.doc/.wps 请先用 convert 转 docx）", file=sys.stderr)
        return 2

    out_dir = Path(args.out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    author = args.author or _system_user()

    # 预检：输入自带修订时 python-docx 读到的正文不可信（实测坑）
    blocked = []
    for path in (a, b):
        ins, dele = _tracked_counts(path)
        if ins or dele:
            blocked.append((path, ins, dele))
    if blocked:
        for path, ins, dele in blocked:
            print(
                f"⚠️ 警告: {Path(path).name} 自身带修订标记（w:ins={ins} w:del={dele}），"
                "python-docx 读出的正文会「删除已生效、插入不存在」，文本 diff 不可信。",
                file=sys.stderr,
            )
        if not args.allow_tracked:
            print(
                "错误: 已拒绝比对。请先另存为不含修订的版本（或用 WPS 接受/拒绝全部修订），"
                "确认可接受不准确的文本 diff 时加 --allow-tracked 继续。",
                file=sys.stderr,
            )
            return 3

    a_lines = _docx_units(a)
    b_lines = _docx_units(b)
    diff = _write_text_diff(a, b, out_dir, a_lines, b_lines)
    print(f"OK: 文本差异 {len(diff)} 行 -> {out_dir / 'diff.txt'}")
    print(f"OK: 可视化对照 -> {out_dir / 'diff.html'}")

    # 实测坑：WPS CompareDocuments 不把页眉/页脚差异记入 Revisions，接受/拒绝都无法还原页眉
    def _part_units(lines):
        return [x for x in lines if x.startswith(("[页眉", "[页脚"))]

    if _part_units(a_lines) != _part_units(b_lines):
        print(
            "⚠️ 警告: 检测到页眉/页脚差异——WPS 原生比较不会把它记入 Revisions，"
            "接受/拒绝修订都无法还原页眉，红线修订只覆盖正文与表格；页眉差异请看 diff.txt。",
            file=sys.stderr,
        )

    tracked = (out_dir / "tracked.docx").resolve()
    try:
        count, summary = _wps_compare(a, b, str(tracked))
    except Exception as exc:
        print(f"错误: WPS 原生比较失败（{type(exc).__name__}: {exc}）", file=sys.stderr)
        return 4

    ins, dele = _tracked_counts(str(tracked))
    if ins + dele <= 0:
        print(
            f"错误: 修订版已生成但包内没有 w:ins/w:del 标记（Revisions={count}）——疑似空产物，请重跑。",
            file=sys.stderr,
        )
        return 5
    fixed = _rewrite_revision_author(tracked, author)
    print(
        f"OK: 红线修订 -> {tracked}（作者={author} Revisions={count} "
        f"w:ins={ins} w:del={dele} 作者改写={fixed} 处）"
    )
    # 文件里的作者已统一改写为 author（COM 读到的原值是 WPS 显示名，故这里按改写后显示）
    for idx, rtype, _rauthor, text in summary[:8]:
        print(f"    修订[{idx}] type={rtype} author={author} text={text!r}")
    return 0


def _print_roles(rep, prefix=""):
    """打印角色识别报告（可回溯）：各角色计数 + 原文层级纠正 + 样本。"""
    r = rep.get("roles")
    if not r:
        return
    counts = r.get("counts") or {}
    order = ["doc_title", "heading_1", "heading_2", "heading_3", "heading_4", "body", "toc"]
    parts = [f"{k}={counts[k]}" for k in order if k in counts]
    parts += [f"{k}={counts[k]}" for k in sorted(counts) if k not in order]
    print(f"{prefix}角色识别: " + "、".join(parts) + f"（赋/改命名样式 {r.get('styled', 0)} 段）")
    corrected = r.get("corrected") or []
    if corrected:
        print(f"{prefix}⚠️  纠正原文层级标记 {len(corrected)} 处（以编号模式为准）：")
        for c in corrected[:10]:
            print(f"{prefix}    段{c['para']}: {c['via']}　「{c['text']}」")
    for s in (r.get("samples") or [])[:8]:
        print(f"{prefix}    样本 段{s['para']} → {s['role']}（{s['via']}）「{s['text']}」")


def cmd_apply_style(args):
    """对已有 docx 套版式（页面 + 命名样式 + 字体四属性）；内容零改动断言不通过则拒绝产出。"""
    import style_spec
    import word_style

    spec = style_spec.load_spec(args.spec)
    style_spec.check_spec_supported(spec)
    src = Path(args.file)
    print(style_spec.describe_spec(spec))
    if args.dry_run:
        rep = word_style.apply_word_style(src, spec, None, dry_run=True)
        print(f"[dry-run] 将套命名样式 {len(rep['styles'])} 个（{', '.join(rep['styles'])}）、"
              f"页面 {rep['page']} 节、清理 run 级字体 {rep['runs_stripped']} 处；未写盘。")
        _print_roles(rep, "  ")
        return 0
    try:
        rep, final, bak = style_spec.commit_style(
            src, args.out, lambda tmp: word_style.apply_word_style(src, spec, tmp), "style")
    except style_spec.ContentChangedError as exc:
        print(f"错误: {exc}", file=sys.stderr)
        print("提示: 已拒绝产出，原文件未改动。请检查该文档是否有工具无法安全保留的结构。", file=sys.stderr)
        return style_spec.ContentChangedError.exit_code
    print(f"✅ 已套版式: {final}")
    print(f"   命名样式 {len(rep['styles'])} 个（{', '.join(rep['styles'])}）｜页面 {rep['page']} 节｜"
          f"清理 run 级字体 {rep['runs_stripped']} 处")
    _print_roles(rep, "   ")
    if bak:
        print(f"   原文件备份: {bak}")
    print("   内容零改动断言: 通过（逐段 + 逐表格单元格比对，差异 0）")
    return 0


def cmd_table_style(args):
    """表格样式：表头底纹/加粗/居中 + 边框 + 跨页重复表头 + 表内字号。"""
    import style_spec
    import word_style

    spec = style_spec.load_spec(args.spec)
    style_spec.check_spec_supported(spec)
    src = Path(args.file)
    print(style_spec.describe_spec(spec))
    if args.dry_run:
        rep = word_style.apply_word_table_style(src, spec, None, dry_run=True)
        print(f"[dry-run] 将处理 {rep['tables']} 张表（表头 {rep['header_cells']} 格 / 正文 {rep['body_cells']} 格）；未写盘。")
        return 0
    try:
        rep, final, bak = style_spec.commit_style(
            src, args.out, lambda tmp: word_style.apply_word_table_style(src, spec, tmp), "tstyle")
    except style_spec.ContentChangedError as exc:
        print(f"错误: {exc}", file=sys.stderr)
        print("提示: 已拒绝产出，原文件未改动。", file=sys.stderr)
        return style_spec.ContentChangedError.exit_code
    print(f"✅ 已套表格样式: {final}")
    print(f"   表格 {rep['tables']} 张（表头 {rep['header_cells']} 格 / 正文 {rep['body_cells']} 格）")
    if bak:
        print(f"   原文件备份: {bak}")
    print("   内容零改动断言: 通过（逐段 + 逐表格单元格比对，差异 0）")
    return 0


def main():
    parser = argparse.ArgumentParser(description="Word 文档工具")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("read", help="提取全文")
    p.add_argument("file")
    p.set_defaults(fn=cmd_read)

    p = sub.add_parser("info", help="统计信息")
    p.add_argument("file")
    p.set_defaults(fn=cmd_info)

    p = sub.add_parser("create", help="从 Markdown 生成 docx")
    p.add_argument("out")
    p.add_argument("--title")
    p.add_argument("--from-md")
    p.add_argument("--body")
    p.set_defaults(fn=cmd_create)

    p = sub.add_parser("edit", help="查找替换文本")
    p.add_argument("file")
    p.add_argument("--replace", action="append", required=True, metavar="旧=新")
    p.add_argument("--out")
    p.set_defaults(fn=cmd_edit)

    p = sub.add_parser("convert", help="转换格式（WPS COM）")
    p.add_argument("src")
    p.add_argument("dst")
    p.set_defaults(fn=cmd_convert)

    p = sub.add_parser("compare", help="文档比对：文本 diff + HTML 对照 + 红线修订（WPS COM）")
    p.add_argument("doc_a", metavar="A.docx")
    p.add_argument("doc_b", metavar="B.docx")
    p.add_argument("--out-dir", required=True, help="输出目录（diff.txt / diff.html / tracked.docx）")
    p.add_argument("--author", default=None, help="修订作者名（默认系统用户名）")
    p.add_argument("--allow-tracked", action="store_true", help="输入含修订标记时仍继续（默认拒绝）")
    p.set_defaults(fn=cmd_compare)

    p = sub.add_parser("apply-style", help="对已有 docx 套版式（页面/命名样式/字体四属性；内容零改动断言）")
    p.add_argument("file")
    p.add_argument("--spec", default=None, help="规格 id 或 JSON 路径（默认 standard）")
    p.add_argument("--out", help="输出文件（默认就地改 + 备份 .bak-style-<时间戳>）")
    p.add_argument("--dry-run", action="store_true", help="只报告不写盘")
    p.set_defaults(fn=cmd_apply_style)

    p = sub.add_parser("table-style", help="表格样式（表头底纹/加粗/居中、边框、跨页重复表头；内容零改动）")
    p.add_argument("file")
    p.add_argument("--spec", default=None, help="规格 id 或 JSON 路径（默认 standard）")
    p.add_argument("--out")
    p.add_argument("--dry-run", action="store_true")
    p.set_defaults(fn=cmd_table_style)

    args = parser.parse_args()
    cli_guard.check_inputs(args)
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(cli_guard.run(main))
