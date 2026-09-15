# -*- coding: utf-8 -*-
"""Excel 套样式（apply-style）—— A 线样式能力实现。

红线：
  1. **只改格式，绝不改值/公式**（34 号教训：任何文本改动都是事故）；
  2. 落盘前后**逐工作表逐单元格比对 value**（含公式原文），差异 → ContentChangedError；
  3. 零新增依赖（openpyxl + 标准库）。
已知边界：openpyxl 保存会丢弃图表 / 图片等「富部件」，本模块会在检测到时**主动告警**（不静默）。
"""
from __future__ import annotations

import zipfile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from doc_roles import resolve_column_align
from style_spec import assert_content_unchanged, excel_snapshot

_RICH_PARTS = ("xl/charts/", "xl/media/", "xl/drawings/")


def detect_rich_parts(path) -> list:
    """检测 openpyxl 无法保真的富部件（图表/图片/绘图）。"""
    found = []
    try:
        with zipfile.ZipFile(str(path)) as zf:
            names = zf.namelist()
    except Exception:  # noqa: BLE001
        return found
    for prefix in _RICH_PARTS:
        if any(n.startswith(prefix) for n in names):
            found.append(prefix.rstrip("/"))
    return found


def _text_width(value) -> int:
    s = str(value)
    return sum(2 if ord(ch) > 127 else 1 for ch in s)


def _match_rules(header: str, rules: list) -> dict:
    if not header:
        return {}
    key = str(header)
    for rule in rules or []:
        for kw in rule.get("header_contains", []) or []:
            if kw and kw in key:
                return rule
    return {}


def apply_excel_style(src, spec: dict, out=None, dry_run: bool = False, sheets=None) -> dict:
    xspec = spec.get("excel") or {}
    if not xspec:
        raise ValueError(f"规格 {spec.get('id')} 不含 excel 段")

    before = excel_snapshot(src)
    rich = detect_rich_parts(src)
    wb = load_workbook(str(src), data_only=False)
    fcfg = xspec.get("font") or {}
    font_name = str(fcfg.get("name", "宋体"))
    font_size = float(fcfg.get("size", 11))
    hdr = xspec.get("header") or {}
    hdr_row = int(hdr.get("row", 1))
    border_cfg = xspec.get("border") or {}
    thick = Side(style=str(border_cfg.get("style", "thin")), color=str(border_cfg.get("color", "808080")))
    border = Border(left=thick, right=thick, top=thick, bottom=thick)
    rules = xspec.get("column_rules") or []
    widths = xspec.get("column_width") or {}
    wrap_cfg = xspec.get("wrap") or {}
    report = {"sheets": [], "rich_parts": rich, "cells": 0}

    for ws in wb.worksheets:
        if sheets and ws.title not in sheets:
            continue
        max_row, max_col = ws.max_row or 0, ws.max_column or 0
        if not max_row or not max_col:
            continue
        matched = {}
        headers = {}
        for ci in range(1, max_col + 1):
            cell = ws.cell(row=hdr_row, column=ci)
            headers[ci] = cell.value
            rule = _match_rules(cell.value, rules)
            if rule:
                matched[ci] = rule

        # 列对齐（A2.1）：序号列居中 / 数值列右对齐 / 文本列按 text_align_default
        xar = xspec.get("align_rules") or {}
        col_align = {}
        if xar:
            for ci in range(1, max_col + 1):
                vals = [ws.cell(row=ri, column=ci).value for ri in range(hdr_row + 1, max_row + 1)]
                col_align[ci] = resolve_column_align(headers.get(ci), vals, {"align_rules": xar})

        for ri in range(1, max_row + 1):
            for ci in range(1, max_col + 1):
                cell = ws.cell(row=ri, column=ci)
                if isinstance(cell, MergedCell):
                    continue
                old = cell.font
                is_header = ri == hdr_row
                cell.font = Font(
                    name=font_name,
                    size=font_size,
                    bold=True if (is_header and hdr.get("bold", True)) else old.bold,
                    italic=old.italic,
                    color=old.color,
                )
                if is_header:
                    if hdr.get("shading"):
                        cell.fill = PatternFill("solid", fgColor=str(hdr["shading"]))
                    cell.alignment = Alignment(
                        horizontal=str(hdr.get("align", "center")),
                        vertical=str(hdr.get("valign", "center")),
                        wrap_text=True,
                    )
                else:
                    rule = matched.get(ci, {})
                    wrap = bool(rule.get("wrap"))
                    if not wrap and wrap_cfg.get("auto_for_long_text"):
                        wrap = _text_width(cell.value) > int(wrap_cfg.get("max_width_chars", 60))
                    cell.alignment = Alignment(horizontal=col_align.get(ci), vertical="center", wrap_text=wrap)
                    if rule.get("number_format"):
                        cell.number_format = str(rule["number_format"])
                if border_cfg:
                    cell.border = border
                report["cells"] += 1

        # 列宽：规则优先，其余按内容自适应
        for ci in range(1, max_col + 1):
            rule = matched.get(ci, {})
            letter = get_column_letter(ci)
            if rule.get("width"):
                ws.column_dimensions[letter].width = float(rule["width"])
            elif widths.get("auto", True):
                best = float(widths.get("min", 6))
                top = float(widths.get("max", 60))
                for ri in range(1, max_row + 1):
                    v = ws.cell(row=ri, column=ci).value
                    if v is None:
                        continue
                    best = min(max(best, _text_width(v) + 2), top)
                ws.column_dimensions[letter].width = best

        if hdr.get("freeze"):
            ws.freeze_panes = str(hdr["freeze"])

        pr = xspec.get("print") or {}
        if pr:
            ws.page_setup.orientation = str(pr.get("orientation", "portrait"))
            if str(pr.get("paper", "A4")).upper() == "A4":
                ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            ws.page_setup.fitToWidth = int(pr.get("fit_to_width", 1))
            ws.page_setup.fitToHeight = 0
            ws.print_options.gridLines = bool(pr.get("gridlines", False))

        report["sheets"].append({"sheet": ws.title, "rows": max_row, "cols": max_col, "ruled_cols": len(matched)})

    if dry_run:
        wb.close()
        return report

    target = str(out or src)
    wb.save(target)
    wb.close()
    assert_content_unchanged(before, excel_snapshot(target), "excel")
    return report
