#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PPT 存量美化（apply-style）—— B 线施工 ④。

对**已有** pptx 套用主题规格：统一文本字体（latin / ea / cs 三属性显式设置，防中文回落），
可选统一文字色；落盘前后做**内容零改动断言**（逐页文本 / 表格单元格 / 图表系列类别 / 备注），
任一差异 → **exit 3 拒绝产出、原文件不动**（机制与 A 线 word/excel 的 apply-style 同构）。

命令：
  py -3 ppt_style.py apply-style <file.pptx> [--spec standard] [--out o.pptx] [--dry-run]
                                 [--text-color ROLE]

退出码：0 成功 ｜ 2 输入/参数/规格错 ｜ 3 **内容零改动断言失败（已拒绝产出、原文件未动）**
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent
for _p in (str(HERE), str(SCRIPTS)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import cli_guard      # noqa: E402
import style_spec     # noqa: E402

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

_HEX6 = re.compile(r"^[0-9A-Fa-f]{6}$")


def _resolve_color(spec, token):
    """--text-color 的取值解析：6 位 hex，或规格 color_roles / 顶层 colors 的键。"""
    t = str(token).strip()
    if _HEX6.match(t):
        return t.upper()
    roles = (spec.get("pptx") or {}).get("color_roles") or {}
    value = roles.get(t)
    if isinstance(value, str):
        if _HEX6.match(value):
            return value.upper()
        node = spec.get("colors") or {}
        for part in value.split("."):
            node = node.get(part) if isinstance(node, dict) else None
        if isinstance(node, str) and _HEX6.match(node):
            return node.upper()
    raise cli_guard.InputError(
        "--text-color 无法解析：%s（可用 6 位 hex，或规格 pptx.color_roles 的键，如 text_on_light）" % token)


def _set_run_fonts(run, latin, ea):
    """显式设置 run 的 a:latin / a:ea / a:cs（中文字形必须落在 a:ea 上，否则会被主题替换掉）。"""
    from pptx.oxml.ns import qn

    rPr = run.font._rPr
    for tag, name in (("a:latin", latin), ("a:ea", ea), ("a:cs", ea)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def _style_text_frame(tf, latin, ea, color_rgb, counters):
    for para in tf.paragraphs:
        for run in para.runs:
            if not run.text or not run.text.strip():
                continue                      # 空 run 不动（避免凭空造出格式）
            _set_run_fonts(run, latin, ea)
            counters["runs"] += 1
            if color_rgb is not None:
                color = run.font.color
                if color is None or color.type is None:      # 只改「未显式设色」的 run，不破坏既有配色
                    run.font.color.rgb = color_rgb
                    counters["colored"] += 1


def _style_shape(shape, latin, ea, color_rgb, counters):
    if shape.has_text_frame:
        _style_text_frame(shape.text_frame, latin, ea, color_rgb, counters)
        counters["shapes"] += 1
    if getattr(shape, "has_table", False):
        try:
            for row in shape.table.rows:
                for cell in row.cells:
                    _style_text_frame(cell.text_frame, latin, ea, color_rgb, counters)
            counters["tables"] += 1
        except Exception:
            pass


def apply_ppt_style(src, spec, out=None, dry_run=False, text_color=None) -> dict:
    """对已有 pptx 统一字体（+ 可选文字色）。dry_run 只报告不写盘；写盘后立即做零改动断言。"""
    from pptx import Presentation
    from pptx.dml.color import RGBColor

    ppspec = spec.get("pptx") or {}
    if not ppspec:
        raise cli_guard.InputError("规格 %s 不含 pptx 段，无法用于 PPT 存量美化" % spec.get("id"))
    fonts = ppspec.get("fonts") or {}
    body = fonts.get("body") or fonts.get("heading") or {}
    latin = str(body.get("latin") or "微软雅黑")
    ea = str(body.get("ea") or latin)
    allowed = [str(x) for x in (ppspec.get("allowed_fonts") or [])]
    if allowed and (latin not in allowed or ea not in allowed):
        raise cli_guard.InputError("要套用的字体 %s / %s 不在规格 allowed_fonts 白名单内（%s）"
                                   % (latin, ea, "、".join(allowed)))

    color_rgb = RGBColor.from_string(_resolve_color(spec, text_color)) if text_color else None

    src = Path(src)
    before = style_spec.pptx_snapshot(src)
    prs = Presentation(str(src))
    counters = {"slides": 0, "shapes": 0, "runs": 0, "tables": 0, "notes": 0, "colored": 0}
    for slide in prs.slides:
        counters["slides"] += 1
        for shape in slide.shapes:
            _style_shape(shape, latin, ea, color_rgb, counters)
        if slide.has_notes_slide:
            try:
                _style_text_frame(slide.notes_slide.notes_text_frame, latin, ea, color_rgb, counters)
                counters["notes"] += 1
            except Exception:
                pass
    report = {"spec": str(spec.get("id")), "latin": latin, "ea": ea,
              "text_color": str(text_color or ""), "dry_run": bool(dry_run)}
    report.update(counters)
    if dry_run:
        return report

    target = Path(out) if out else src
    prs.save(str(target))
    # 内容零改动断言：任一差异 → ContentChangedError（调用方转 exit 3，拒绝产出）
    style_spec.assert_content_unchanged(before, style_spec.pptx_snapshot(target), "pptx")
    return report


def _print_report(rep, prefix=""):
    print("%sslides %d ｜ 文本框 %d ｜ run %d ｜ 表格 %d ｜ 备注 %d ｜ 文字色处理 %d"
          % (prefix, rep["slides"], rep["shapes"], rep["runs"], rep["tables"], rep["notes"], rep["colored"]))
    print("   字体 latin / ea：%s / %s%s"
          % (rep["latin"], rep["ea"], ("；文字色：%s" % rep["text_color"]) if rep["text_color"] else ""))


def cmd_apply_style(args):
    spec = style_spec.load_spec(args.spec, for_format="ppt")
    src = Path(args.file)
    print(style_spec.describe_spec(spec))
    if args.dry_run:
        rep = apply_ppt_style(src, spec, None, dry_run=True, text_color=args.text_color)
        _print_report(rep, "[dry-run] ")
        print("（--dry-run：未写文件）")
        return 0
    try:
        rep, final, bak = style_spec.commit_style(
            src, args.out,
            lambda tmp: apply_ppt_style(src, spec, tmp, text_color=args.text_color), "pptstyle")
    except style_spec.ContentChangedError as exc:
        print("错误: %s" % exc, file=sys.stderr)
        print("提示: 内容零改动断言未通过 —— 已拒绝产出、原文件未改动（exit 3）。"
              "常见原因：该文件含工具无法安全保留的结构（组合图形内文本、SmartArt、图表标题等）。", file=sys.stderr)
        return style_spec.ContentChangedError.exit_code
    print("OK: 已套样式 %s" % final)
    _print_report(rep)
    if bak:
        print("   原文件备份: %s" % bak)
    return 0


def main():
    parser = argparse.ArgumentParser(description="PPT 存量美化（apply-style）")
    sub = parser.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("apply-style", help="统一已有 pptx 的字体（+ 可选文字色）；内容零改动断言，不通过则 exit 3")
    p.add_argument("file", help="输入 .pptx")
    p.add_argument("--spec", default="standard", help="主题规格 id 或路径（默认 standard）")
    p.add_argument("--out", help="另存路径；省略则就地修改并自动备份")
    p.add_argument("--dry-run", action="store_true", help="只报告将改什么，不写文件")
    p.add_argument("--text-color", dest="text_color",
                   help="统一文字色（6 位 hex 或规格 color_roles 的键，如 text_on_light）；"
                        "只作用于未显式设色的 run，默认关")
    p.set_defaults(fn=cmd_apply_style)
    args = parser.parse_args()
    cli_guard.check_inputs(args)
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(cli_guard.run(main))
