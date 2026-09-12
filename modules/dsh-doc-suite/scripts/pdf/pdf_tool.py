"""PDF 工具（PyMuPDF / pdfplumber / pypdf）。

用法（Windows 一律用 py -3，`python` 可能是 Microsoft Store 别名 stub）：
  py -3 pdf_tool.py info <pdf>
  py -3 pdf_tool.py text <pdf> [--pages 1-3,5] [--out out.txt]
  py -3 pdf_tool.py tables <pdf> [--pages ..] [--out out.xlsx|csv|json]
  py -3 pdf_tool.py merge <out.pdf> a.pdf b.pdf ...
  py -3 pdf_tool.py split <pdf> --pages 1-3 --out part.pdf
  py -3 pdf_tool.py images <pdf> [--pages ..] --outdir dir [--dpi 150]
  py -3 pdf_tool.py make <out.pdf> img1.png img2.jpg ... [--fit a4|auto]
  py -3 pdf_tool.py ocr <pdf>                  # 【已退役】仅打印退役说明，不执行 OCR

参数形态（实测易踩）：**输出参数在前**——`merge <out> <files...>`、`make <out> <图片...>`；
`split --pages` 与 `images --outdir` 必填。

说明：
  - text/tables 只对"文字型 PDF"有效；扫描件请先 images 转图交给基座原生识图（本机不做 OCR）。
  - text 会告警：旋转页（page.rotation != 0）→ 表格列序/阅读顺序可能反转。
  - text 会判定扫描件：整页无文本层且含图片时明确回报"此页为扫描件"，并给出转图片建议。
  - tables 走 PyMuPDF find_tables()，输出附 bbox 坐标；检测到合并单元格/空单元格时告警"该表可能失真"。
    找不到表格**不是错误**（exit 0），提示走 stderr。
  - info 会校验文件头是否为 %PDF，非 PDF 只给尽力解析结果并告警。
  - ocr 通道已于 2026-09-12 **退役**：不再本地 OCR、也不再调用云端视觉 API。
"""
import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import cli_guard  # noqa: E402

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")


def parse_pages(spec, total):
    """'1-3,5' -> [1,2,3,5]（1 起始）。None 表示全部。"""
    if not spec:
        return list(range(1, total + 1))
    pages = []
    for part in spec.split(","):
        part = part.strip()
        m = re.match(r"^(\d+)(?:-(\d+))?$", part)
        if not m:
            raise ValueError(f"无效页码: {part}")
        start = int(m.group(1))
        end = int(m.group(2) or start)
        if end < start:
            raise ValueError(f"无效页码范围: {part}")
        pages.extend(range(start, end + 1))
    return [p for p in pages if 1 <= p <= total]


def _looks_like_pdf(path):
    """文件头判定：前 1KB 里应出现 %PDF 魔术字（少数 PDF 前置垃圾字节）。"""
    try:
        with open(path, "rb") as fh:
            head = fh.read(1024)
    except OSError:
        return False
    return b"%PDF" in head


def cmd_info(args):
    import pymupdf as fitz

    if not _looks_like_pdf(args.pdf):
        print(
            "⚠️ 警告: 该文件未检测到 %PDF 文件头，**它可能不是 PDF**"
            "（扩展名或内容不符）；以下结果由 PyMuPDF 尽力解析，仅供参考。",
            file=sys.stderr,
        )
    doc = fitz.open(args.pdf)
    print(f"文件: {args.pdf}")
    print(f"页数: {doc.page_count}")
    print(f"元数据: {json.dumps(doc.metadata, ensure_ascii=False)}")
    doc.close()
    return 0


def _scan_info(page):
    """扫描件判定：整页几乎无文本层且含图片。返回 (是否扫描件, 文本字符数, 图片数, 图片覆盖比)。"""
    try:
        text = page.get_text("text") or ""
    except Exception:
        text = ""
    images = page.get_images(full=True) or []
    try:
        page_area = abs(page.rect.width * page.rect.height) or 1.0
        covered = 0.0
        for info in page.get_image_info():
            x0, y0, x1, y1 = info.get("bbox", (0, 0, 0, 0))
            covered += max(0.0, x1 - x0) * max(0.0, y1 - y0)
        ratio = min(1.0, covered / page_area)
    except Exception:
        ratio = 0.0
    stripped = text.strip()
    is_scan = len(stripped) < 20 and len(images) > 0
    return is_scan, len(stripped), len(images), ratio


def _warn_rotation(page, n):
    rotation = getattr(page, "rotation", 0)
    if rotation:
        print(
            f"⚠️ 警告: 第 {n} 页 rotation={rotation}（旋转页）——表格列序与阅读顺序可能反转，"
            "需人工核对或转图片后核对。",
            file=sys.stderr,
        )
    return rotation


def cmd_text(args):
    import pymupdf as fitz

    doc = fitz.open(args.pdf)
    pages = parse_pages(args.pages, doc.page_count)
    parts = []
    for n in pages:
        page = doc.load_page(n - 1)
        _warn_rotation(page, n)
        is_scan, chars, img_count, ratio = _scan_info(page)
        text = page.get_text("text")
        if is_scan:
            print(
                f"⚠️ 警告: 第 {n} 页为扫描件——无文本层（{chars} 字符，图片 {img_count} 张，覆盖约 {ratio:.0%}）。"
                f"建议：转图片交给基座原生识图（py -3 scripts\\pdf\\pdf_tool.py images \"{args.pdf}\" "
                f"--pages {n} --outdir <目录>），本机不做本地 OCR。",
                file=sys.stderr,
            )
            if not text.strip():
                text = (
                    f"[此页为扫描件：无文本层，图片 {img_count} 张，覆盖约 {ratio:.0%}；"
                    "本机不做本地 OCR，建议转图片交给基座原生识图]\n"
                )
        parts.append(f"--- 第 {n} 页 ---\n{text}" if len(pages) > 1 else text)
    doc.close()
    out = "\n".join(parts)
    if args.out:
        Path(args.out).write_text(out, encoding="utf-8")
        print(f"OK: 已写出 {len(out)} 字符 -> {args.out}")
    else:
        print(out)
    return 0


def _merged_cell_suspect(rows):
    """启发式：表头或表体出现空单元格 -> 疑似合并单元格/跨列表头，提取结果可能失真。"""
    if not rows or len(rows) < 2:
        return False
    header_blank = any(not c for c in rows[0]) and len(rows[0]) > 1
    body_blank = any(not c for row in rows[1:] for c in row)
    return bool(header_blank or body_blank)


def _tables_pdfplumber(path, n):
    """兜底：pdfplumber 提表（同样带 bbox）。"""
    import pdfplumber

    found = []
    try:
        with pdfplumber.open(path) as pdf:
            page = pdf.pages[n - 1]
            for table in page.find_tables():
                rows = [[("" if c is None else str(c).strip()) for c in row] for row in table.extract()]
                found.append({
                    "page": n,
                    "bbox": [round(float(v), 1) for v in table.bbox],
                    "rows": rows,
                    "engine": "pdfplumber",
                })
    except Exception as exc:
        print(f"⚠️ 警告: 第 {n} 页 pdfplumber 兜底也失败（{type(exc).__name__}: {exc}）", file=sys.stderr)
    return found


def _find_tables_quiet(page):
    """调 page.find_tables() 并吞掉 PyMuPDF 往 stdout 打的库提示，返回 [(bbox, rows)]。"""
    import contextlib
    import io

    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        finder = page.find_tables()
        found = []
        for table in finder.tables:
            rows = [[("" if c is None else str(c).strip()) for c in row] for row in table.extract()]
            found.append(([round(float(v), 1) for v in table.bbox], rows))
    return found


def cmd_tables(args):
    import pymupdf as fitz

    doc = fitz.open(args.pdf)
    pages = parse_pages(args.pages, doc.page_count)
    tables = []
    for n in pages:
        page = doc.load_page(n - 1)
        _warn_rotation(page, n)
        found = []
        try:
            for bbox, rows in _find_tables_quiet(page):
                found.append({"page": n, "bbox": bbox, "rows": rows, "engine": "pymupdf"})
        except Exception as exc:
            print(f"⚠️ 警告: 第 {n} 页 PyMuPDF find_tables 失败（{type(exc).__name__}: {exc}），回退 pdfplumber。", file=sys.stderr)
        if not found:
            found = _tables_pdfplumber(args.pdf, n)
        for table in found:
            if _merged_cell_suspect(table["rows"]):
                box = ", ".join(str(v) for v in table["bbox"])
                print(
                    f"⚠️ 警告: 第 {n} 页表格 bbox=({box}) 疑似含合并单元格/跨列表头（存在空单元格），"
                    "该表可能失真，需人工复核。",
                    file=sys.stderr,
                )
        tables.extend(found)
    doc.close()

    if not tables:
        is_scan, chars, img_count, ratio = (False, 0, 0, 0.0)
        probe = fitz.open(args.pdf)
        if probe.page_count:
            is_scan, chars, img_count, ratio = _scan_info(probe.load_page(pages[0] - 1) if pages else probe.load_page(0))
        probe.close()
        if is_scan:
            print(
                f"⚠️ 警告: 该页为扫描件（{chars} 字符，图片 {img_count} 张，覆盖约 {ratio:.0%}），无文本层可提表；"
                "建议转图片交给基座原生识图（py -3 scripts\\pdf\\pdf_tool.py images ...），本机不做本地 OCR。",
                file=sys.stderr,
            )
        else:
            print(
                "提示: 未提取到表格（可能是无边框表格、纯图片排版或扫描件）——这不是错误。"
                "若是扫描件，请用 images 转图后交基座原生识图。",
                file=sys.stderr,
            )
        return 0

    suffix = Path(args.out).suffix.lower() if args.out else ""
    if suffix == ".json":
        Path(args.out).write_text(json.dumps(tables, ensure_ascii=False, indent=2), encoding="utf-8")
    elif suffix == ".csv":
        with open(args.out, "w", encoding="utf-8-sig", newline="") as fh:
            import csv

            writer = csv.writer(fh)
            for t in tables:
                writer.writerow([f"[第 {t['page']} 页]"])
                writer.writerow([f"# bbox=({', '.join(str(v) for v in t['bbox'])}) engine={t['engine']}"])
                writer.writerows(t["rows"])
    elif suffix == ".xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        for ti, t in enumerate(tables, 1):
            ws = wb.create_sheet(f"表{ti}-p{t['page']}")
            ws.append([f"bbox=({', '.join(str(v) for v in t['bbox'])}) engine={t['engine']}"])
            for row in t["rows"]:
                ws.append(row)
        wb.save(args.out)
    else:
        for t in tables:
            box = ", ".join(str(v) for v in t["bbox"])
            print(f"== 第 {t['page']} 页 ==  表 bbox=({box}) 引擎={t['engine']}")
            for row in t["rows"]:
                print(" | ".join(row))
        return 0
    print(f"OK: {len(tables)} 张表格 -> {args.out}")
    return 0


def cmd_merge(args):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for f in args.files:
        writer.append(f)
    writer.write(args.out)
    print(f"OK: 合并 {len(args.files)} 个文件 -> {args.out}")
    return 0


def cmd_split(args):
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(args.pdf)
    pages = parse_pages(args.pages, len(reader.pages))
    writer = PdfWriter()
    for n in pages:
        writer.add_page(reader.pages[n - 1])
    writer.write(args.out)
    print(f"OK: 提取 {len(pages)} 页 -> {args.out}")
    return 0


def cmd_images(args):
    import pymupdf as fitz

    doc = fitz.open(args.pdf)
    pages = parse_pages(args.pages, doc.page_count)
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    for n in pages:
        pix = doc.load_page(n - 1).get_pixmap(dpi=args.dpi)
        out = outdir / f"page-{n:03d}.png"
        pix.save(out)
        print(f"OK: {out}")
    doc.close()
    return 0


def cmd_make(args):
    import pymupdf as fitz

    doc = fitz.open()
    for f in args.images:
        img = fitz.open(f)
        rect = img[0].rect
        if args.fit == "a4":
            page = doc.new_page(width=595, height=842)
            scale = min(595 / rect.width, 842 / rect.height) * 0.95
            w, h = rect.width * scale, rect.height * scale
            page.insert_image(fitz.Rect((595 - w) / 2, (842 - h) / 2, (595 + w) / 2, (842 + h) / 2), filename=f)
        else:
            page = doc.new_page(width=rect.width, height=rect.height)
            page.insert_image(rect, filename=f)
        img.close()
    doc.save(args.out)
    print(f"OK: {len(args.images)} 张图片 -> {args.out}")
    return 0


OCR_RETIRED_NOTE = """OCR 通道已于 2026-09-12 退役——本工具不再做本地 OCR，也不再调用任何云端视觉 API。

替代做法（推荐顺序）：
  1. 直接把图片/扫描件发给基座（DSH 原生识图），无需本工具参与；
  2. 只想把扫描页取出来核对：py -3 pdf_tool.py images <pdf> --outdir <目录>
  3. 判断 PDF 是文字型还是扫描件：py -3 pdf_tool.py text <pdf>（会明确回报"此页为扫描件"）
"""


def cmd_ocr(args):
    """墓碑：保留子命令以免旧脚本静默跑偏，但只打印退役说明（exit 0）。"""
    print(OCR_RETIRED_NOTE, file=sys.stderr)
    return 0


def main():
    parser = argparse.ArgumentParser(description="PDF 工具")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("info")
    p.add_argument("pdf")
    p.set_defaults(fn=cmd_info)

    p = sub.add_parser("text")
    p.add_argument("pdf")
    p.add_argument("--pages")
    p.add_argument("--out")
    p.set_defaults(fn=cmd_text)

    p = sub.add_parser("tables")
    p.add_argument("pdf")
    p.add_argument("--pages")
    p.add_argument("--out")
    p.set_defaults(fn=cmd_tables)

    p = sub.add_parser("merge")
    p.add_argument("out")
    p.add_argument("files", nargs="+")
    p.set_defaults(fn=cmd_merge)

    p = sub.add_parser("split")
    p.add_argument("pdf")
    p.add_argument("--pages", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(fn=cmd_split)

    p = sub.add_parser("images")
    p.add_argument("pdf")
    p.add_argument("--pages")
    p.add_argument("--outdir", required=True)
    p.add_argument("--dpi", type=int, default=150)
    p.set_defaults(fn=cmd_images)

    p = sub.add_parser("make")
    p.add_argument("out")
    p.add_argument("images", nargs="+")
    p.add_argument("--fit", choices=["a4", "auto"], default="auto")
    p.set_defaults(fn=cmd_make)

    p = sub.add_parser("ocr", help="【已退役】原 OCR 通道；请改用基座原生识图",
                       description="【已退役·2026-09-12】OCR 通道不再可用：本工具不做本地 OCR，"
                                   "也不调用任何云端视觉 API。请把扫描件/图片直接交给基座原生识图。")
    p.add_argument("pdf")
    p.add_argument("--pages")
    p.add_argument("--out")
    p.add_argument("--dpi", type=int, default=200)
    p.set_defaults(fn=cmd_ocr)

    args = parser.parse_args()
    cli_guard.check_inputs(args)
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(cli_guard.run(main))
